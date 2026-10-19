from pathlib import Path
import csv
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
root=Path(__file__).resolve().parent;assets=root/'report-images';assets.mkdir(exist_ok=True)
rows=list(csv.DictReader((root/'results/intervals.csv').open()))
for r in rows:
    for key in ('start','end'):r[key]=float(r[key])
base=min(r['start'] for r in rows);end=max(r['end'] for r in rows);span=end-base
font=lambda n:ImageFont.truetype('C:/Windows/Fonts/arial.ttf',n)
def canvas(title):
    im=Image.new('RGB',(1600,1000),'white');d=ImageDraw.Draw(im);d.text((40,15),title,font=font(32),fill='black')
    for i in range(9):
        x=200+i*165;d.line((x,90,x,890),fill='#dddddd');d.text((x-20,915),f'{span*i/8/1000:.1f}',font=font(22),fill='black')
    d.text((700,955),'Время от первого обращения, с',font=font(26),fill='black');return im,d
def x(t):return 200+(t-base)/span*1320
im,d=canvas('Состояния 8 читателей и 8 писателей')
actors=[(role,str(i)) for role in ('reader','writer') for i in range(1,9)]
for j,(role,i) in enumerate(actors):
    y=100+j*48;d.text((10,y),f'{role} {i}',font=font(24),fill='black');prev=base
    for r in sorted([r for r in rows if r['role']==role and r['id']==i],key=lambda r:r['start']):
        d.rectangle((x(prev),y+5,x(r['start']),y+30),fill='#eeeeee');d.rectangle((x(r['start']),y+5,x(r['end']),y+30),fill='#c75b40' if role=='writer' else '#377db2');d.text((x(r['start'])+3,y+5),str(r['page']),font=font(18),fill='white');d.line((x(r['end']),y,x(r['end']),y+35),fill='black',width=2);prev=r['end']
d.text((210,60),'Синий чтение   красный запись   серый ожидание   чёрная отметка освобождение; число — страница',font=font(21),fill='black');im.save(assets/'states.png')
im,d=canvas('Занятость 16 страниц — число активных обращений')
peak=1
for p in range(16):
    y=100+p*48;d.text((12,y),f'Страница {p}',font=font(23),fill='black')
    for k in range(264):
        t=base+span*(k+.5)/264;n=sum(r['page']==str(p) and r['start']<=t<r['end'] for r in rows);peak=max(peak,n)
        color=['#eeeeee','#80b5dc','#367aaf','#18466b'][min(n,3)];d.rectangle((200+k*5,y+4,205+k*5,y+32),fill=color)
d.text((210,60),'Светло-серый 0   голубой 1   синий 2   тёмно-синий 3 и более',font=font(23),fill='black');im.save(assets/'heatmap.png')
doc=Document(root.parent/'лаба3/Отчет_лабораторная_3.docx');body=doc._element.body
start=next(p._p for p in doc.paragraphs if p.text=='Цель работы')
for el in list(body)[list(body).index(start):]:
    if el.tag!=qn('w:sectPr'):body.remove(el)
for p in doc.paragraphs:
    for r in p.runs:r.text=r.text.replace('РАБОТА № 3','РАБОТА № 4').replace('Процессы и потоки','Межпроцессное взаимодействие')
for line in (root/'REPORT.md').read_text(encoding='utf-8').splitlines():
    if not line or line.startswith('# Работа'):continue
    if line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
    else:
        p=doc.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.first_line_indent=Cm(1.25);p.paragraph_format.line_spacing=1.5
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(14)
doc.add_page_break();doc.add_paragraph('ГРАФИКИ СОСТОЯНИЙ И ЗАНЯТОСТИ',style='Heading 1')
for i,name in enumerate(('states','heatmap'),1):
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(assets/(name+'.png')),width=Cm(16))
    p=doc.add_paragraph(f'Рисунок {i} — '+('Смена состояний процессов' if i==1 else 'Занятость страниц во времени'),style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.save(root/'Отчет_лабораторная_4.docx');print('Word report and charts generated; peak occupancy',peak)
