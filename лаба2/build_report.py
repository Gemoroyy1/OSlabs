from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

root=Path(__file__).resolve().parent
assets=root/'report-images';assets.mkdir(exist_ok=True)
doc=Document(root.parent/'лаба1/Отчет_лабораторная_1.docx')
body=doc._element.body
start=next(p._p for p in doc.paragraphs if p.text=='ЦЕЛЬ РАБОТЫ')
for el in list(body)[list(body).index(start):]:
    if el.tag!=qn('w:sectPr'):body.remove(el)
for p in doc.paragraphs:
    for r in p.runs:
        r.text=r.text.replace('РАБОТА № 1','РАБОТА № 2').replace('Управление файловой системой','Управление памятью')
num=0
def para(text):
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent=Cm(1.25)
    p.paragraph_format.line_spacing=1.5
    p.paragraph_format.space_after=Pt(6)
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(14)
    return p
def heading(text,level=1):
    p=doc.add_paragraph(text,style=f'Heading {level}')
    for r in p.runs:r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
def figure(name,lines,caption):
    global num
    num+=1
    # Reproduce saved output as a labelled protocol illustration.
    font=ImageFont.truetype('C:/Windows/Fonts/consola.ttf',24)
    im=Image.new('RGB',(1600,50+35*len(lines)),(25,29,35));d=ImageDraw.Draw(im)
    for i,line in enumerate(lines):d.text((20,20+35*i),line,font=font,fill='white')
    image=assets/name;im.save(image)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(image),width=Cm(16))
    p=doc.add_paragraph(f'Рисунок {num} — {caption}',style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11)

heading('ЦЕЛЬ РАБОТЫ')
para('Исследовать механизмы управления виртуальной памятью Win32 и обмен данными между процессами через проецируемый файл. Реализовать приложения на C++17, проверить выделение и защиту памяти, а также корректность обмена между клиентом и сервером.')
heading('ЗАДАНИЯ НА ЛАБОРАТОРНУЮ РАБОТУ')
para('Задание 2.1: создать консольное меню для получения информации о системе и памяти, исследования региона по адресу, раздельного и совместного резервирования и фиксации страниц, записи данных, изменения защиты и освобождения памяти.')
para('Задание 2.2: создать Linux-сервер и клиент с меню. Сервер создаёт файл, отображает его, записывает сообщение и ждёт подтверждения чтения. Клиент открывает файл, отображает его, ждёт готовности, читает и выводит данные. При завершении отображения отменяются, сервер удаляет файл.')
heading('СРЕДА ИСПЫТАНИЙ')
para('Задание 2.1 выполнено в Windows x64 с MinGW g++. Задание 2.2 выполнено в Ubuntu 24.04 под WSL2 с g++. Использован стандарт C++17. Иллюстрации ниже воспроизводят сохранённые протоколы программ; они не являются снимками экрана.')
doc.add_page_break()
lines=(root/'REPORT.md').read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line or line.startswith('# Работа') or line=='## Цель':continue
    if line.startswith('Исследовать виртуальное'):continue
    if line.startswith('## '):
        title=line[3:]
        if title in ('Задание 2.2','Выводы'):doc.add_page_break()
        heading(title.upper())
    elif line.startswith('### '):heading(line[4:],2)
    elif line.startswith('|'):
        rows=[line]
        while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
        rows=[r for r in rows if not set(r.replace('|','').replace('-','').strip())==set()]
        table=doc.add_table(rows=0,cols=2)
        # Reuse source table formatting without carrying its contents.
        table._tbl.replace(table._tbl.tblPr,deepcopy(doc.tables[0]._tbl.tblPr))
        for row in rows:
            cells=table.add_row().cells
            for c,value in zip(cells,row.strip('|').split('|')):
                c.text=value.strip()
                for p in c.paragraphs:
                    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11)
        log=(root/'windows-tests.txt').read_text(encoding='utf-8-sig').splitlines()
        figure('memory-states.png',[x for x in log if 'State=' in x or 'written' in x or 'Previous=' in x or 'error' in x],'Проверка состояний памяти и защиты, выдержка из протокола')
    else:
        para(line)
        if line.startswith('GlobalMemoryStatus используется'):
            log=(root/'windows-tests.txt').read_text(encoding='utf-8-sig').splitlines()
            figure('system-memory.png',[x for x in log if 'Page=' in x or 'Load=' in x],'Параметры системы и статус памяти, выдержка из протокола')
        if line.startswith('Ubuntu 24.04, WSL2'):
            log=(root/'linux-tests.txt').read_text(encoding='utf-8').splitlines()
            figure('exchange.png',[x for x in log if x.startswith('CASE') or 'received:' in x or 'acknowledged' in x],'Три успешных обмена между процессами, выдержка из протокола')
heading('ПРИЛОЖЕНИЕ')
para('memory_win.cpp — управление памятью Windows. shared.cpp — сервер и клиент Linux. test_exchange.py — автоматическая проверка обмена. windows-tests.txt, windows-tests-additional.txt и linux-tests.txt — первичные протоколы запусков. Команды сборки и порядок работы приведены в README.md.')
out=root/'Отчет_лабораторная_2.docx';doc.save(out)
print(out)
