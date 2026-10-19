from pathlib import Path
import csv,statistics
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
root=Path(__file__).resolve().parent
rows=list(csv.DictReader((root/'results/measurements.csv').open()))
ns=[1,2,4,8,12,16];stats={}
for impl in ('win32','openmp'):
    for n in ns:
        values=[float(r['seconds'])*1000 for r in rows if r['implementation']==impl and int(r['threads'])==n]
        stats[impl,n]=(statistics.median(values),min(values),max(values))
text='''# Работа 3 Процессы и потоки

## Цель работы
Исследовать создание, распределение работы и завершение потоков средствами Win32 API и OpenMP. Вычислить π методом средних прямоугольников и сравнить время выполнения.

## Задание и исходные данные
Формула: π ≈ (1/N) Σ 4/(1+xᵢ²), xᵢ=(i+0,5)/N, i=0…N−1. N=10000000 — число итераций, а не число гарантированных верных знаков. Студенческий билет 230616, размер блока B=10×230616=2306160. Реализованы оба приложения на C++17 для Windows.

Всего пять блоков: четыре полных и последний из 775360 итераций. При n>5 дополнительные потоки не получают вычислительных заданий, однако создаются и корректно завершаются. Уже при n=4 один поток получает второй блок, поэтому балансировка не идеальна.

## Задание 3.1 Win32 API
Пул создаётся в начале программы через CreateThread с CREATE_SUSPENDED. Первые блоки выдаются по очереди. Рабочий поток вычисляет частичную сумму, записывает её в свой слот и сигнализирует событие завершения. Диспетчер ожидает события через WaitForMultipleObjects, приостанавливает завершивший блок поток через SuspendThread, забирает сумму и выдаёт следующий свободный блок через ResumeThread.

Потоки не завершаются между блоками. После окончания всех заданий устанавливается атомарный флаг stop, все потоки возобновляются и завершаются кооперативно. WaitForSingleObject подтверждает завершение; события и потоки закрываются через CloseHandle. Время измеряется QueryPerformanceCounter.

Отклонение от буквального самоприостановления: SuspendThread вызывает диспетчер после события, а не сам рабочий поток. Это устраняет гонку раннего ResumeThread перед фактическим самоприостановлением. Между заданиями рабочий поток проверяет только атомарные флаги и уступает процессор SwitchToThread. SuspendThread используется здесь по требованию учебной работы; в прикладной программе предпочтительнее ожидание событий без принудительной остановки потоков.

## Задание 3.2 OpenMP
Работа распределяется директивой for с schedule(dynamic,2306160). Каждый освободившийся поток получает очередной блок. Частичные суммы объединяются reduction(+:sum), что исключает гонку при обновлении суммы. Число потоков задаётся num_threads, omp_set_dynamic(0) отключает динамическое сокращение команды. Фактическое число потоков проверено во всех запусках. Время измеряется steady_clock.

## Методика проверки
Среда: Windows x64, Intel Core i7-11800H. Сборка MinGW g++ с -std=c++17 -O2 -Wall -Wextra; для OpenMP дополнительно -fopenmp. Для каждой реализации выполнены 10 повторов при n=1,2,4,8,12,16 — всего 120 запусков. Порядок перемешан с фиксированным seed=230616; перед измерениями сделан прогрев каждой реализации.

Проверены успешное завершение, размер блока, число фактически созданных потоков и абсолютная ошибка π относительно math.pi менее 10⁻¹¹. Все проверки успешны. Обычный запуск Win32 с n=4 дал 3,1415926535897776; OpenMP — 3,1415926535897678. Небольшие различия объясняются порядком сложения чисел double.

Таймер охватывает создание команды, вычисление и завершение потоков; запуск процесса и вывод CSV находятся вне измеряемого интервала. Привязка к ядрам и фиксация частоты процессора не выполнялись. Из-за короткого времени запуска важны накладные расходы и фоновые процессы. Основной показатель — медиана десяти повторов; минимум и максимум отражают разброс.

## Результаты измерений
| Реализация | Потоки | Медиана мс | Минимум мс | Максимум мс | Ускорение |
|---|---|---|---|---|---|
'''
for impl in ('win32','openmp'):
    for n in ns:
        med,lo,hi=stats[impl,n]
        text+=f'| {impl} | {n} | {med:.3f} | {lo:.3f} | {hi:.3f} | {stats[impl,1][0]/med:.2f} |\n'
text+='\n## Выводы\n'
for impl in ('win32','openmp'):
    best=min(ns,key=lambda n:stats[impl,n][0]);text+=f'Для {impl} лучший результат среди испытанных вариантов: {best} потоков, медиана {stats[impl,best][0]:.3f} мс.\n\n'
text+='''Динамическое распределение позволяет передать освободившемуся потоку следующий блок. Однако здесь доступно только пять блоков, поэтому увеличение числа потоков сверх пяти не увеличивает полезный параллелизм. Результаты ограничены конкретной средой и крупным блоком, заданным номером студбилета; максимум медианы не является универсальным оптимумом.

OpenMP скрывает управление потоками и объединение сумм, тогда как Win32 требует явного протокола событий, приостановления и завершения. Обе реализации дают согласованный результат π. В отличие от второй лабораторной, где проверялись виртуальная память и межпроцессное отображение файла, здесь исследуется вычислительный параллелизм. Сравнивать время этих разных задач как показатель ускорения некорректно; количественное сравнение выполнено между Win32 и OpenMP для одной формулы.

## Приложение
Исходники: pi_win32.cpp и pi_openmp.cpp. benchmark.py воспроизводит измерения. Все 120 исходных записей находятся в results/measurements.csv; среда и параметры — в results/environment.json. Размер блока не изменялся для улучшения результатов.
'''
(root/'REPORT.md').write_text(text,encoding='utf-8')
assets=root/'report-images';assets.mkdir(exist_ok=True)
def chart(speed=False):
    im=Image.new('RGB',(1500,720),'white');d=ImageDraw.Draw(im);font=lambda n:ImageFont.truetype('C:/Windows/Fonts/arial.ttf',n)
    d.text((80,20),'Ускорение относительно n=1' if speed else 'Время вычисления π, мс, медиана 10 повторов',font=font(34),fill='black')
    vals={impl:[stats[impl,1][0]/stats[impl,n][0] if speed else stats[impl,n][0] for n in ns] for impl in ('win32','openmp')};maximum=max(max(v) for v in vals.values())*1.15
    for k in range(6):
        y=590-k*94;d.line((100,y,1430,y),fill='#dddddd');d.text((8,y-15),f'{maximum*k/5:.1f}',font=font(24),fill='black')
    for impl,color in [('win32','#205b9e'),('openmp','#b34228')]:
        coords=[(150+i*245,590-v/maximum*470) for i,v in enumerate(vals[impl])];d.line(coords,fill=color,width=5)
        for (x,y),v in zip(coords,vals[impl]):d.ellipse((x-6,y-6,x+6,y+6),fill=color);d.text((x-30,y-30 if impl=='win32' else y+8),f'{v:.2f}',font=font(21),fill=color)
    for i,n in enumerate(ns):d.text((140+i*245,610),str(n),font=font(26),fill='black')
    d.text((500,665),'Число потоков',font=font(27),fill='black');d.text((1030,75),'Win32',font=font(26),fill='#205b9e');d.text((1230,75),'OpenMP',font=font(26),fill='#b34228')
    path=assets/('speedup.png' if speed else 'time.png');im.save(path);return path
doc=Document(root.parent/'лаба2/Отчет_лабораторная_2.docx');body=doc._element.body;start=next(p._p for p in doc.paragraphs if p.text=='ЦЕЛЬ РАБОТЫ')
for el in list(body)[list(body).index(start):]:
    if el.tag!=qn('w:sectPr'):body.remove(el)
for p in doc.paragraphs:
    for r in p.runs:r.text=r.text.replace('РАБОТА № 2','РАБОТА № 3').replace('Управление памятью','Процессы и потоки')
for line in text.splitlines():
    if not line or line.startswith('# Работа') or line.startswith('|---'):continue
    if line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('|'):
        # Render compact measured table as actual Word table.
        if 'Реализация' in line:t=doc.add_table(rows=0,cols=6)
        cells=t.add_row().cells
        for c,s in zip(cells,line.strip('|').split('|')):
            c.text=s.strip()
            for p in c.paragraphs:
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)
    else:
        p=doc.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.line_spacing=1.5;p.paragraph_format.first_line_indent=Cm(1.25)
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(14)
doc.add_page_break();doc.add_paragraph('ГРАФИКИ РЕЗУЛЬТАТОВ',style='Heading 1')
for i,speed in enumerate((False,True),1):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(chart(speed)),width=Cm(16))
    p=doc.add_paragraph(f'Рисунок {i} — '+('Ускорение' if speed else 'Время выполнения'),style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.save(root/'Отчет_лабораторная_3.docx')
print(text[text.index('## Выводы'):text.index('OpenMP скрывает')])
